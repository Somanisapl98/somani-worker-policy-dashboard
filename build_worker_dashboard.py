from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date
from html import escape
import json, re, sys

BASE=Path(__file__).resolve().parent
INPUT=BASE/'input'; OUTPUT=BASE/'output'; HISTORY=BASE/'history'; LOGS=BASE/'logs'
for d in (INPUT,OUTPUT,HISTORY,LOGS): d.mkdir(exist_ok=True)
POLICY_PREFIXES=('IONIRI','IONIRA','IONIKO','IONIBI','ISOMRA','IONIAM','IONIJA','ISOMRI','ISOMIL','ISOMOA','ISOMAM')
EXCLUDE=('SAPL','CTTRY','CTRY')
TODAY=date.today()

def log(msg):
    line=f"{datetime.now().isoformat(timespec='seconds')} | {msg}"
    print(line)
    with (LOGS/'worker_update.log').open('a',encoding='utf-8') as f:f.write(line+'\n')

def norm(s): return re.sub(r'[^A-Z0-9]','',str(s or '').upper())
def dtval(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    for fmt in ('%d-%m-%Y','%d/%m/%Y','%Y-%m-%d','%d-%b-%y','%d-%b-%Y'):
        try:return datetime.strptime(str(v).strip(),fmt).date()
        except:pass
    return None

def newest(pattern, folder=INPUT):
    files=[p for p in folder.glob(pattern) if p.is_file() and not p.name.startswith('~$')]
    return max(files,key=lambda p:p.stat().st_mtime) if files else None

def parse_amount(cell):
    raw=cell.value if cell is not None else None
    text=str(raw or '').strip(); upper=text.upper(); side=''
    if re.search(r'\bDR\s*$',upper):side='DR'
    elif re.search(r'\bCR\s*$',upper):side='CR'
    cleaned=re.sub(r'(?i)\s*(DR|CR)\s*$','',text).replace(',','').strip()
    try: amt=float(cleaned) if cleaned else float(raw or 0)
    except:return 0.0,''
    fmt=str(getattr(cell,'number_format','') or '').upper()
    if not side and ('DR' in fmt or 'CR' in fmt):
        parts=fmt.split(';'); section=parts[0 if amt>=0 else min(1,len(parts)-1)]
        side='DR' if 'DR' in section else ('CR' if 'CR' in section else '')
    if not side:side='CR' if amt<0 else 'DR'
    return abs(amt),side

def customer_type(name, code, mintify_ids):
    return 'Mintify' if norm(code) in mintify_ids else 'Non-Mintify'

def slab(age, typ):
    if age is None:return None
    if typ=='Mintify':
        return 1 if age<=25 else 2 if age<=45 else 3 if age<=70 else 4 if age<=100 else 5
    return 1 if age<=25 else 2 if age<=70 else 3 if age<=100 else 4

def load_mintify():
    folder=INPUT/'mintify'; folder.mkdir(exist_ok=True)
    files=[p for p in folder.glob('*.xlsx') if not p.name.startswith('~$')]
    if not files:return set()
    p=max(files,key=lambda x:x.stat().st_mtime); w=load_workbook(p,read_only=True,data_only=True)
    ids=set()
    for ws in w.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                n=norm(v)
                if n and any(c.isdigit() for c in n): ids.add(n)
    return ids

def main():
    src=newest('OT DT*.xlsx')
    if not src: raise SystemExit('ERROR: No OT DT*.xlsx found in input folder.')
    log(f'Using source: {src.name}')
    ws=load_workbook(src,read_only=True,data_only=True).worksheets[0]
    hr=None; headers=[]
    for rn,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,30),values_only=True),1):
        vals=[str(v or '').strip().lower().replace('’',"'") for v in row]
        if 'date' in vals and any('ref. no' in v or 'ref no' in v for v in vals) and any("party's name" in v for v in vals):hr=rn;headers=vals;break
    if not hr: raise SystemExit('ERROR: Tally header row not found.')
    def ci(*terms):
        for i,h in enumerate(headers):
            if any(t in h for t in terms): return i
        return None
    di,ri,pi,pendi,codei=ci('date'),ci('ref. no','ref no'),ci("party's name"),ci('pending'),ci("party's code",'party code')
    if None in (di,ri,pi,pendi):raise SystemExit('ERROR: Required Tally columns are missing.')
    party=code=''; records=[]
    for row in ws.iter_rows(min_row=hr+1,values_only=False):
        val=lambda i: row[i].value if i is not None and i<len(row) else None
        dt,ref,pname=val(di),str(val(ri) or '').strip(),str(val(pi) or '').strip()
        if not dt and not ref and pname:
            m=re.match(r'^(.*)\(([^()]*)\)\s*$',pname); party=m.group(1).strip() if m else pname; code=m.group(2).strip() if m else code; continue
        if pname:
            m=re.match(r'^(.*)\(([^()]*)\)\s*$',pname); party=m.group(1).strip() if m else pname
            if codei is not None and val(codei) not in (None,''):code=str(val(codei)).strip()
            elif m:code=m.group(2).strip()
        invu=ref.upper()
        if any(x in invu for x in EXCLUDE):continue
        d=dtval(dt)
        if not(d and ref and party):continue
        amt,side=parse_amount(row[pendi])
        prefix=next((p for p in POLICY_PREFIXES if invu.startswith(p)), '')
        if not(prefix and side=='DR' and amt>0):continue
        records.append({'party':party,'code':code,'age':max(0,(TODAY-d).days),'prefix':prefix})
    mintify=load_mintify(); parties={}
    for r in records:
        z=parties.setdefault(r['party'],{'codes':set(),'oldest':0})
        if r['code']:z['codes'].add(str(r['code']))
        z['oldest']=max(z['oldest'],r['age'])
    current={}
    for name,z in parties.items():
        typ=customer_type(name,next(iter(z['codes']),''),mintify)
        current[name]={'type':typ,'slab':slab(z['oldest'],typ),'ids':sorted(z['codes'])}
    snapfile=HISTORY/'worker_policy_snapshot.json'
    previous=json.loads(snapfile.read_text(encoding='utf-8')) if snapfile.exists() else {}
    same_source_file=HISTORY/'worker_source.json'
    previous_source=json.loads(same_source_file.read_text(encoding='utf-8')).get('source','') if same_source_file.exists() else ''
    same_source=(previous_source==src.name)
    rows=[]
    for name,z in current.items():
        old=previous.get(name,{}).get('slab')
        cur=z['slab']
        if same_source: movement='No new comparison'
        elif old is None: movement='New party'
        elif old==cur: movement=f'No change (Slab {cur})'
        elif cur>old: movement=f'Slab {old} → Slab {cur}'
        else: movement=f'Slab {old} → Slab {cur}'
        rows.append({'party':name,'type':z['type'],'movement':movement,'previous':old,'current':cur,'ids':z['ids']})
    if not same_source:
        for name,old in previous.items():
            if name not in current: rows.append({'party':name,'type':old.get('type',''),'movement':'No policy bills','previous':old.get('slab'),'current':None,'ids':old.get('ids',[])})
        snapfile.write_text(json.dumps(current,indent=2),encoding='utf-8')
        same_source_file.write_text(json.dumps({'source':src.name,'processed_at':datetime.now().isoformat(timespec='seconds')},indent=2),encoding='utf-8')
    rows.sort(key=lambda x:(str(x['movement']),x['party']))
    data=json.dumps(rows,ensure_ascii=False).replace('</','<\\/')
    html=f'''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>SOMANI Worker Policy Dashboard</title><style>
*{{box-sizing:border-box}}body{{margin:0;font-family:Segoe UI,Arial;background:#f5f7fb;color:#10213a}}header{{background:#123f78;color:white;padding:22px 28px;border-bottom:7px solid #f2b21b}}header h1{{margin:0;font-size:27px}}header p{{margin:5px 0 0;font-size:13px}}main{{padding:16px}}.notice{{background:#eaf2fc;border-left:5px solid #123f78;padding:10px 14px;margin-bottom:12px;border-radius:6px}}.filters{{display:grid;grid-template-columns:2fr repeat(3,1fr);gap:10px;background:white;padding:14px;border-radius:12px;border:1px solid #d7e0ec;position:sticky;top:0;z-index:3}}label{{display:block;font-size:12px;font-weight:700;margin-bottom:5px}}input,select{{width:100%;padding:10px;border:1px solid #b9c7d8;border-radius:7px;background:white}}.list{{margin-top:12px;display:flex;flex-direction:column;gap:8px}}.row{{display:grid;grid-template-columns:minmax(260px,2fr) 1fr 1.3fr 1fr 1fr;gap:12px;align-items:center;background:white;border:1px solid #d7e0ec;border-radius:9px;padding:11px 13px}}.party{{font-weight:750}}.party small{{display:block;font-weight:500;color:#526174;margin-top:3px}}.cell b{{display:block;font-size:10px;text-transform:uppercase;color:#64748b;margin-bottom:3px}}.movement{{font-weight:750;padding:6px 9px;background:#fff3c7;border-radius:14px;text-align:center}}.worse{{background:#fee2e2;color:#991b1b}}.better{{background:#dcfce7;color:#166534}}.none{{background:#e5e7eb;color:#374151}}.empty{{background:white;padding:30px;text-align:center;border-radius:10px}}@media(max-width:850px){{.filters{{grid-template-columns:1fr 1fr}}.row{{grid-template-columns:1fr 1fr}}.party{{grid-column:1/-1}}}}@media(max-width:520px){{.filters{{grid-template-columns:1fr}}.row{{grid-template-columns:1fr}}}}
</style></head><body><header><h1>SOMANI Worker Policy Dashboard</h1><p>Policy movement only | Source: {escape(src.name)} | Updated: {datetime.now().strftime('%d %b %Y %I:%M %p')}</p></header><main><div class="notice"><b>Worker view:</b> Monetary values, totals, invoice amounts, KPI figures, invoice details and export options are intentionally not included. Policy logic uses the same 11 prefixes and open DR bills only. {'Same source, no new daily comparison.' if same_source else 'New source comparison.'}</div><div class="filters"><div><label>Party or ID</label><input id="q" placeholder="Search party name or ID"></div><div><label>Customer type</label><select id="type"><option>All</option><option>Mintify</option><option>Non-Mintify</option></select></div><div><label>Movement</label><select id="movement"><option>All</option></select></div><div><label>Current slab</label><select id="slab"><option value="All">All slabs</option><option value="1">Slab 1</option><option value="2">Slab 2</option><option value="3">Slab 3</option><option value="4">Slab 4</option><option value="5">Slab 5</option><option value="None">No policy bill</option></select></div></div><div class="list" id="list"></div></main><script>const DATA={data};const $=id=>document.getElementById(id),esc=s=>String(s??'').replace(/[&<>"']/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[c]));function cls(m){{if(m.includes('No new'))return'none';const x=m.match(/Slab (\d+) → Slab (\d+)/);if(!x)return'';return +x[2]>+x[1]?'worse':'better'}}function render(){{let q=$('q').value.trim().toUpperCase(),t=$('type').value,m=$('movement').value,s=$('slab').value;let a=DATA.filter(x=>(!q||(x.party+' '+x.ids.join(' ')).toUpperCase().includes(q))&&(t==='All'||x.type===t)&&(m==='All'||x.movement===m)&&(s==='All'||(s==='None'?x.current==null:String(x.current)===s)));$('list').innerHTML=a.length?a.map(x=>`<div class="row"><div class="party">${{esc(x.party)}}<small>${{esc(x.type)}} | ID: ${{esc(x.ids.join(' / ')||'Not available')}}</small></div><div class="cell"><b>Previous slab</b>${{x.previous==null?'Not available':'Slab '+x.previous}}</div><div class="movement ${{cls(x.movement)}}">${{esc(x.movement)}}</div><div class="cell"><b>Current slab</b>${{x.current==null?'No policy bill':'Slab '+x.current}}</div><div class="cell"><b>Status</b>${{x.current&&((x.type==='Mintify'&&x.current===5)||(x.type==='Non-Mintify'&&x.current===4))?'No Billing':'Billing allowed'}}</div></div>`).join(''):'<div class="empty">No matching parties</div>'}}[...new Set(DATA.map(x=>x.movement))].sort().forEach(v=>{{let o=document.createElement('option');o.textContent=o.value=v;$('movement').appendChild(o)}});['q','type','movement','slab'].forEach(id=>$ (id).addEventListener(id==='q'?'input':'change',render));render();</script></body></html>'''
    (OUTPUT/'index.html').write_text(html,encoding='utf-8')
    log(f'Worker dashboard created: {len(rows)} movement records')

if __name__=='__main__':main()
